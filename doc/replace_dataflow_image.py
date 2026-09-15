#!/usr/bin/env python3
"""
Replace the data-flow diagram picture on the already-inserted diagram slide
with the regenerated image, re-fitting it (aspect preserved, no overflow).
Identifies the target slide by its title text.
"""

from pathlib import Path
from PIL import Image
from pptx import Presentation
from pptx.util import Inches

DOC_DIR = Path(__file__).parent
MERMAID_DIR = DOC_DIR / "mermaid"

TARGETS = {
    "pipecat-voicebot-voiceapi-architecture-JP.pptx": ("図解: データフロー", "data_flow_jp.png"),
    "pipecat-voicebot-voiceapi-architecture.pptx": ("Diagram: Data Flow", "data_flow_en.png"),
}


def slide_title(slide):
    for shape in slide.shapes:
        if shape.has_text_frame and shape.text_frame.text.strip():
            return shape.text_frame.text.strip()
    return ""


def fit_rect(prs, img_path):
    slide_w, slide_h = prs.slide_width, prs.slide_height
    margin = Inches(0.4)
    top_offset = Inches(1.15)
    bottom_margin = Inches(0.35)
    avail_w = slide_w - 2 * margin
    avail_h = slide_h - top_offset - bottom_margin

    with Image.open(img_path) as im:
        aspect = im.width / im.height

    if aspect > avail_w / avail_h:
        draw_w = avail_w
        draw_h = int(draw_w / aspect)
    else:
        draw_h = avail_h
        draw_w = int(draw_h * aspect)

    left = margin + (avail_w - draw_w) // 2
    top = top_offset + (avail_h - draw_h) // 2
    return left, top, draw_w, draw_h


def main():
    for pptx_name, (title, png) in TARGETS.items():
        pptx_path = DOC_DIR / pptx_name
        img_path = MERMAID_DIR / png
        prs = Presentation(str(pptx_path))

        target_slide = None
        for slide in prs.slides:
            if slide_title(slide) == title:
                target_slide = slide
                break

        if target_slide is None:
            print(f"⚠️  Slide '{title}' not found in {pptx_name}")
            continue

        # Remove existing picture(s)
        pics = [s for s in target_slide.shapes if s.shape_type == 13]
        for pic in pics:
            pic._element.getparent().remove(pic._element)

        left, top, w, h = fit_rect(prs, img_path)
        target_slide.shapes.add_picture(str(img_path), left, top, width=w, height=h)

        prs.save(str(pptx_path))
        print(f"✓ Replaced data-flow image in {pptx_name} "
              f"(size {w/914400:.2f}x{h/914400:.2f} in, pos {left/914400:.2f},{top/914400:.2f})")


if __name__ == "__main__":
    main()
