#!/usr/bin/env python3
"""
Insert Mermaid diagram images into PowerPoint presentations.
Images are scaled preserving aspect ratio, centered, and never overflow the slide.
Each diagram is placed on its own new slide inserted right after the relevant
existing content slide.
"""

import copy
from pathlib import Path
from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor

DOC_DIR = Path(__file__).parent
MERMAID_DIR = DOC_DIR / "mermaid"

# Slide geometry (EMU). Standard 10 x 7.5 inch.
EMU_PER_INCH = 914400

# Configuration per language:
# Each entry: (diagram png filename, slide title, insert-after title substring)
CONFIG = {
    "pipecat-voicebot-voiceapi-architecture-JP.pptx": {
        "diagrams": [
            # (png, new slide title, exact title of slide to insert AFTER)
            ("system_architecture_jp.png", "図解: システムアーキテクチャ全体構成", "システムアーキテクチャ"),
            ("pipeline_jp.png", "図解: Pipecat パイプライン詳細", "Pipecat パイプラインフロー"),
            ("data_flow_jp.png", "図解: データフロー", "データフロー: 電話 (PSTN)"),
        ],
    },
    "pipecat-voicebot-voiceapi-architecture.pptx": {
        "diagrams": [
            ("system_architecture_en.png", "Diagram: Full System Architecture", "System Architecture"),
            ("pipeline_en.png", "Diagram: Pipecat Pipeline Detail", "Pipecat Pipeline Flow"),
            ("data_flow_en.png", "Diagram: Data Flow", "Data Flow: Phone (PSTN)"),
        ],
    },
}


def get_slide_title_text(slide):
    """Extract the first meaningful text (title) from a slide."""
    for shape in slide.shapes:
        if shape.has_text_frame:
            text = shape.text_frame.text.strip()
            if text:
                return text
    return ""


def move_slide_after(prs, slide, target_index):
    """Move a slide (currently last) to position right after target_index."""
    xml_slides = prs.slides._sldIdLst
    slides = list(xml_slides)
    # The newly added slide is the last one
    new_slide_element = slides[-1]
    xml_slides.remove(new_slide_element)
    # Insert after target_index -> position target_index+1
    xml_slides.insert(target_index + 1, new_slide_element)


def add_diagram_slide(prs, title, image_path):
    """Create a blank slide with a title and a centered, aspect-preserved image."""
    slide_layout = prs.slide_layouts[6]  # Blank
    slide = prs.slides.add_slide(slide_layout)

    # Background
    bg = slide.background.fill
    bg.solid()
    bg.fore_color.rgb = RGBColor(255, 255, 255)

    slide_w = prs.slide_width
    slide_h = prs.slide_height

    # Title box
    title_h = Inches(0.7)
    title_box = slide.shapes.add_textbox(Inches(0.4), Inches(0.25), slide_w - Inches(0.8), title_h)
    tf = title_box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(28)
    p.font.bold = True
    p.font.color.rgb = RGBColor(25, 118, 210)
    p.alignment = PP_ALIGN.CENTER

    # Available area for the image (below the title, with margins)
    margin = Inches(0.4)
    top_offset = Inches(1.15)   # below title
    bottom_margin = Inches(0.35)

    avail_w = slide_w - 2 * margin
    avail_h = slide_h - top_offset - bottom_margin

    # Determine native image size and aspect ratio
    with Image.open(image_path) as img:
        img_w_px, img_h_px = img.width, img.height
    aspect = img_w_px / img_h_px

    # Fit within available area preserving aspect ratio
    avail_aspect = avail_w / avail_h
    if aspect > avail_aspect:
        # Width-constrained
        draw_w = avail_w
        draw_h = int(draw_w / aspect)
    else:
        # Height-constrained
        draw_h = avail_h
        draw_w = int(draw_h * aspect)

    # Center within available area
    left = margin + (avail_w - draw_w) // 2
    top = top_offset + (avail_h - draw_h) // 2

    slide.shapes.add_picture(str(image_path), left, top, width=draw_w, height=draw_h)
    return slide


def process_presentation(pptx_name, config):
    pptx_path = DOC_DIR / pptx_name
    print(f"\n{'='*60}")
    print(f"Processing: {pptx_name}")
    print(f"{'='*60}")

    prs = Presentation(str(pptx_path))

    # Insert diagrams. We process one at a time, re-scanning titles each time
    # so that "insert after" can also target a diagram slide we just added.
    for png_name, slide_title, after_substr in config["diagrams"]:
        image_path = MERMAID_DIR / png_name
        if not image_path.exists():
            print(f"⚠️  Image not found: {image_path}")
            continue

        # Find target slide index by exact title match (fallback to substring)
        target_index = None
        for idx, slide in enumerate(prs.slides):
            title_text = get_slide_title_text(slide)
            if title_text == after_substr:
                target_index = idx
                break
        if target_index is None:
            for idx, slide in enumerate(prs.slides):
                title_text = get_slide_title_text(slide)
                if after_substr in title_text:
                    target_index = idx
                    break

        # Create the diagram slide (added at the end)
        add_diagram_slide(prs, slide_title, image_path)

        if target_index is not None:
            move_slide_after(prs, prs.slides[-1], target_index)
            print(f"✓ Inserted '{png_name}' after slide #{target_index+1} ('{after_substr}')")
        else:
            print(f"✓ Appended '{png_name}' at the end (no match for '{after_substr}')")

    prs.save(str(pptx_path))
    print(f"✓ Saved: {pptx_path}  (total slides: {len(prs.slides._sldIdLst)})")


def main():
    for pptx_name, config in CONFIG.items():
        process_presentation(pptx_name, config)
    print("\n" + "="*60)
    print("All diagrams inserted successfully!")
    print("="*60)


if __name__ == "__main__":
    main()
